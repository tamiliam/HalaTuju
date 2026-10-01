"""TD-257 — what a sponsor signs, and what a sponsor is shown, driven through HTTP.

Four routes the TD-219 guard listed as never driven by any test:
  * `sponsor-terms/<pk>/sections/`                       PUT  (AdminSponsorTermsSectionsView)
  * `sponsor-terms/<pk>/sections/<order>/generate-quiz/` POST (AdminSponsorTermsGenerateQuizView)
  * `sponsor-terms/<pk>/import-docx/`                    POST (AdminSponsorTermsImportDocxView)
  * `sponsor/graduation-messages/`                       GET  (SponsorGraduationMessagesView)

The terms are a CONSENT document — what a benefactor accepts — so each test reads back what the
service decided rather than what the caller sent: orders assigned by position, a quiz payload
dropped with its flag, a translation kept only if it marks the same answer, a sub-clause folded
into its parent. The two Gemini seams are mocked at the one boundary the module names
(`sponsor_terms._gemini_generate`); no test reaches a model. The .docx is a REAL Word file built
with python-docx, so `contracts._docx_structure` parses it for real.

The relay is the anonymity boundary in the other direction: student words reaching a funder. The
funded student and the approved message are produced by the services that produce them in life
(`fund_student` → `respond_to_award`; `submit_graduation_message` → `approve_graduation_message`).
"""
import io
import json
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings

from apps.scholarship import in_programme, pool, sponsor_terms
from apps.scholarship.models import Sponsor, SponsorTermsVersion
from apps.scholarship.tests.factories import (
    TEST_JWT_SECRET, authed_client, make_admin, make_cohort, unique_suffix,
)
from apps.scholarship.tests.test_endpoints_disbursements import fund_through_the_product

ST = '/api/v1/admin/scholarship/sponsor-terms/'
QUIZ = {'tag': 'T', 'plain': 'p', 'question': 'q?', 'options': ['a', 'b', 'c'],
        'correct': 1, 'why': 'because'}
DOCX_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def _docx_bytes():
    """A real Word file: a title, two clauses, and a sub-clause under the first."""
    from docx import Document
    doc = Document()
    doc.add_paragraph('Sponsor Terms', style='Title')
    doc.add_paragraph('Your gift', style='Heading 1')
    doc.add_paragraph('A gift is paid into the programme wallet.')
    doc.add_paragraph('Timing', style='Heading 2')
    doc.add_paragraph('Tranches follow the semester calendar.')
    doc.add_paragraph('Your data', style='Heading 1')
    doc.add_paragraph('We never show you who the student is.')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@override_settings(ROOT_URLCONF='halatuju.urls', SUPABASE_JWT_SECRET=TEST_JWT_SECRET)
class _TermsCase(TestCase):
    def setUp(self):
        self.author = make_admin('org_admin', owning_org=make_cohort().owning_organisation)
        self.client = authed_client(self.author)
        self.terms = sponsor_terms.create_version(version=unique_suffix('v-'),
                                                  by_email=self.author.email)

    def _put(self, sections, terms=None, client=None):
        return (client or self.client).put(f'{ST}{(terms or self.terms).id}/sections/',
                                           {'sections': sections}, format='json')

    def _published(self):
        terms = sponsor_terms.create_version(version=unique_suffix('pub-'))
        SponsorTermsVersion.objects.filter(pk=terms.pk).update(status='published')
        terms.refresh_from_db()
        return terms


class TestReplaceSections(_TermsCase):
    def test_the_author_replaces_every_section_and_the_server_assigns_the_order(self):
        r = self._put([
            {'order': 9, 'heading_en': 'Your gift', 'body_en': 'Paid per tranche.',
             'is_quiz_candidate': False, 'quiz_en': QUIZ},
            {'order': 3, 'heading_en': 'Your data', 'body_en': 'Anonymous.',
             'is_quiz_candidate': True, 'quiz_en': QUIZ},
        ])
        self.assertEqual(r.status_code, 200, r.content)
        rows = list(self.terms.sections.order_by('order').values_list(
            'order', 'heading_en', 'quiz_en'))
        # Orders come from POSITION, never the payload; an unflagged section's quiz is DROPPED.
        self.assertEqual(rows, [(1, 'Your gift', {}), (2, 'Your data', QUIZ)])
        self.assertEqual([s['order'] for s in r.json()['sections']], [1, 2])

    def test_a_published_version_is_immutable(self):
        published = self._published()
        r = self._put([{'heading_en': 'Sneaky edit'}], terms=published)
        self.assertEqual((r.status_code, r.json()['error']), (400, 'not_draft'))
        self.assertFalse(published.sections.exists())

    def test_a_reviewer_cannot_author_the_terms(self):
        reviewer = make_admin('reviewer', owning_org=self.author.owning_organisation)
        r = self._put([{'heading_en': 'x'}], client=authed_client(reviewer))
        self.assertEqual(r.status_code, 403)
        self.assertFalse(self.terms.sections.exists())


class TestGenerateQuiz(_TermsCase):
    def setUp(self):
        super().setUp()
        sponsor_terms.replace_sections(self.terms, [
            {'heading_en': 'Your gift', 'body_en': 'Paid per tranche.', 'is_quiz_candidate': True},
            {'heading_en': 'Contact', 'body_en': 'Email us.', 'is_quiz_candidate': False},
        ])

    def _generate(self, order):
        return self.client.post(f'{ST}{self.terms.id}/sections/{order}/generate-quiz/', {},
                                format='json')

    @mock.patch('apps.scholarship.sponsor_terms._gemini_generate')
    def test_a_draft_quiz_is_saved_and_a_translation_marking_a_different_answer_is_dropped(self, gen):
        gen.return_value = json.dumps({
            'en': dict(QUIZ), 'ms': dict(QUIZ, question='s?'),
            'ta': dict(QUIZ, question='k?', correct=2),       # disagrees with en
        })
        r = self._generate(1)
        self.assertEqual(r.status_code, 200, r.content)
        gen.assert_called_once()
        section = self.terms.sections.get(order=1)
        self.assertEqual(section.quiz_en['correct'], 1)
        self.assertEqual(section.quiz_ms['question'], 's?')
        self.assertEqual(section.quiz_ta, {})
        self.assertTrue(section.quiz_generated_model)

    @mock.patch('apps.scholarship.sponsor_terms._gemini_generate')
    def test_a_section_not_flagged_for_a_quiz_spends_no_ai_call(self, gen):
        r = self._generate(2)
        self.assertEqual((r.status_code, r.json()['error']), (400, 'quiz_not_candidate'))
        gen.assert_not_called()

    @mock.patch('apps.scholarship.sponsor_terms._gemini_generate')
    def test_an_order_with_no_section_is_not_found(self, gen):
        self.assertEqual(self._generate(7).status_code, 404)
        gen.assert_not_called()


class TestImportDocx(_TermsCase):
    def _upload(self, terms=None):
        f = SimpleUploadedFile('terms.docx', _docx_bytes(), content_type=DOCX_TYPE)
        return self.client.post(f'{ST}{(terms or self.terms).id}/import-docx/', {'file': f},
                                format='multipart')

    def test_a_word_file_becomes_a_proposal_with_sub_clauses_folded_and_nothing_saved(self):
        r = self._upload()
        self.assertEqual(r.status_code, 200, r.content)
        body = r.json()
        self.assertEqual(body['title'], 'Sponsor Terms')
        self.assertEqual([s['heading_en'] for s in body['sections']], ['Your gift', 'Your data'])
        # The `1.1 Timing` sub-clause lives INSIDE section 1, in its original order.
        self.assertIn('Timing', body['sections'][0]['body_en'])
        self.assertIn('Tranches follow the semester calendar.', body['sections'][0]['body_en'])
        # A proposal only: the author reviews it and PUTs the result.
        self.assertFalse(self.terms.sections.exists())

    def test_only_a_draft_accepts_an_import(self):
        r = self._upload(terms=self._published())
        self.assertEqual((r.status_code, r.json()['error']), (400, 'not_draft'))

    def test_a_request_with_no_file_is_refused(self):
        r = self.client.post(f'{ST}{self.terms.id}/import-docx/', {}, format='multipart')
        self.assertEqual((r.status_code, r.json()['error']), (400, 'no_file'))


@override_settings(ROOT_URLCONF='halatuju.urls', SUPABASE_JWT_SECRET=TEST_JWT_SECRET,
                   SPONSOR_POOL_ENABLED=True)
class TestGraduationRelay(TestCase):
    """GET sponsor/graduation-messages/ — approved thank-yous from the students THIS sponsor funds."""

    def setUp(self):
        cohort = make_cohort()
        reviewer = make_admin('reviewer', owning_org=cohort.owning_organisation)
        self.app, self.sponsor, _sp = fund_through_the_product(cohort, reviewer=reviewer)
        approved = in_programme.submit_graduation_message(
            self.app, raw_text='Thank you for believing in me.')
        in_programme.approve_graduation_message(approved, by_email='staff@example.test')
        in_programme.submit_graduation_message(self.app, raw_text='A second note, not yet read.')

    def _get(self, sponsor):
        return authed_client(sponsor).get('/api/v1/sponsor/graduation-messages/')

    def test_the_funder_reads_only_the_approved_message_under_the_anonymous_ref(self):
        r = self._get(self.sponsor)
        self.assertEqual(r.status_code, 200, r.content)
        messages = r.json()['messages']
        self.assertEqual([(m['ref'], m['text']) for m in messages],
                         [(pool.pool_ref(self.app.id), 'Thank you for believing in me.')])
        self.assertEqual(set(messages[0]), {'ref', 'text', 'approved_at'})

    def test_a_sponsor_who_funds_nobody_reads_nothing(self):
        stranger = Sponsor.objects.create(
            supabase_user_id=unique_suffix('sponsor-uid-'), name='Other Funder',
            email=f'{unique_suffix("other-")}@example.test', status='approved')
        r = self._get(stranger)
        self.assertEqual((r.status_code, r.json()['messages']), (200, []))

    def test_dark_while_the_pool_is_switched_off(self):
        with override_settings(SPONSOR_POOL_ENABLED=False):
            r = self._get(self.sponsor)
        self.assertEqual(r.status_code, 404)
